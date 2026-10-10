#!/usr/bin/env python3
"""Build an editable, evidence-based advisor presentation from deck_content.json.

Requires python-pptx and Pillow. No repository runtime or experiment is executed.
"""
from pathlib import Path
import json, math, sys

if Path('/tmp/agentflow-report-pydeps').exists():
    sys.path.insert(0, '/tmp/agentflow-report-pydeps')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent
DATA = json.loads((ROOT / 'deck_content.json').read_text())
OUT = ROOT / 'AgentFlow_自进化系统_导师汇报_20261009.pptx'
W,H = 13.333333,7.5
NAVY='142B45'; INK='23384E'; MUTED='637587'; TEAL='138F8B'
BLUE='3974BD'; PURPLE='7660AD'; ORANGE='D47E31'; RED='B74747'
BG='F5F7FA'; LINE='DBE3EB'; WHITE='FFFFFF'; PALE='E6F3F1'
FONT='Noto Sans CJK SC'
prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
prs.core_properties.title='AgentFlow 自进化系统：当前机制、实验进展与研究方向'
prs.core_properties.subject='30分钟导师汇报；当前源码与现存实验产物核查'
prs.core_properties.author='AgentFlow research workspace'
prs.core_properties.keywords='AgentFlow, self-evolution, GEPA, AppWorld'

def rgb(v): return RGBColor.from_string(v)

def box(s,x,y,w,h,fill=WHITE,line=None,radius=True):
    z=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                         Inches(x),Inches(y),Inches(w),Inches(h))
    if radius:
        z.adjustments[0]=.12
    z.fill.solid(); z.fill.fore_color.rgb=rgb(fill)
    z.line.fill.background() if line is None else None
    if line is not None: z.line.color.rgb=rgb(line); z.line.width=Pt(.8)
    return z

def text(s,x,y,w,h,content,size=19,color=INK,bold=False,align=PP_ALIGN.LEFT):
    z=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=z.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.01)
    tf.margin_top=tf.margin_bottom=Inches(.01)
    lines=content if isinstance(content,list) else str(content).split('\n')
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(5); p.line_spacing=1.14
        r=p.add_run(); r.text=str(line); r.font.name=FONT
        r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
        rp=r._r.get_or_add_rPr(); ea=OxmlElement('a:ea'); ea.set('typeface',FONT); rp.append(ea)
    return z

def line(s,x1,y1,x2,y2,color=LINE,width=1.5,arrow=False):
    z=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    z.line.color.rgb=rgb(color); z.line.width=Pt(width)
    if arrow:
        end=OxmlElement('a:tailEnd'); end.set('type','triangle'); end.set('w','sm'); end.set('len','sm')
        z._element.spPr.get_or_add_ln().append(end)
    return z

def chip(s,x,y,w,label,color=TEAL):
    box(s,x,y,w,.34,fill=color,radius=True)
    text(s,x+.07,y+.025,w-.14,.28,label,10.5,WHITE,True,PP_ALIGN.CENTER)

def base(d,i):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
    box(s,0,0,W,.075,fill=TEAL,radius=False)
    text(s,.55,.30,10.9,.34,d.get('section',''),11,TEAL,True)
    text(s,.55,.81,12.1,.70,d['title'],27 if len(d['title'])>33 else 29,NAVY,True)
    text(s,.56,1.59,12.05,.61,d.get('subtitle',''),13.5,MUTED)
    line(s,.55,7.07,12.78,7.07)
    footer=d.get('source_short','当前源码与现存产物 · 证据索引见讲稿')
    if len(footer)>120: footer=footer[:110]+'…（详见备注）'
    text(s,.55,7.15,11.5,.23,footer,8.5,MUTED)
    text(s,12.18,7.13,.62,.27,f'{i:02d} / {len(DATA["slides"]):02d}',9.5,MUTED,align=PP_ALIGN.RIGHT)
    return s

def takeaway(s,msg,color=TEAL,y=6.42):
    box(s,.56,y,12.18,.48,fill=PALE if color==TEAL else 'FFF0E3')
    text(s,.72,y+.085,11.85,.31,msg,14,color,True)

def card(s,x,y,w,h,title,body,color=TEAL,index=None,body_size=17):
    box(s,x,y,w,h,line=LINE)
    box(s,x,y,.055,h,fill=color,radius=False)
    if index:
        text(s,x+.21,y+.20,w-.42,.36,index,12,color,True)
        ty=y+.69
    else: ty=y+.23
    text(s,x+.21,ty,w-.42,.65,title,20,NAVY,True)
    text(s,x+.21,ty+.79,w-.42,h-(ty-y)-.94,body,body_size,INK)

def table(s,headers,rows,widths=None,y=2.30,h=3.80,font=15):
    widths=widths or [1/len(headers)]*len(headers)
    x=.56; total=12.18
    rh=h/(len(rows)+1)
    for ri,row in enumerate([headers]+rows):
        cx=x
        for ci,item in enumerate(row):
            cw=widths[ci]*total
            fill=NAVY if ri==0 else (WHITE if ri%2 else 'EAF0F5')
            box(s,cx,y+ri*rh,cw,rh,fill=fill,radius=False)
            text(s,cx+.12,y+ri*rh+.12,cw-.24,rh-.17,str(item),font,WHITE if ri==0 else INK,ri==0)
            cx+=cw
    return s

def render_cards(s,d):
    cards=d['cards']; n=len(cards); cols=2 if n in (2,4,6) else min(n,3)
    rows=math.ceil(n/cols); gap=.22; cw=(12.18-gap*(cols-1))/cols
    height=(3.85-gap*(rows-1))/rows; y=2.25
    for k,c in enumerate(cards):
        card(s,.56+(k%cols)*(cw+gap),y+(k//cols)*(height+gap),cw,height,
             c['title'],c['body'],c.get('color',[TEAL,BLUE,PURPLE,ORANGE][k%4]),c.get('index'))

def render_flow(s,d):
    steps=d['steps']; n=len(steps); cw=(12.18-.26*(n-1))/n; y=2.36
    for k,c in enumerate(steps):
        x=.56+k*(cw+.26)
        card(s,x,y,cw,2.95,c['title'],c['body'],c.get('color',TEAL),f'{k+1:02d}',15)
        if k<n-1: line(s,x+cw+.01,y+1.25,x+cw+.25,y+1.25,TEAL,2,True)
    if d.get('below'): text(s,.70,5.59,11.9,.64,d['below'],16,MUTED)

def render_lanes(s,d):
    for k,lane in enumerate(d['lanes']):
        y=2.32+k*1.78; color=TEAL if k==0 else BLUE
        box(s,.56,y,12.18,1.50,fill=WHITE,line=LINE)
        chip(s,.76,y+.16,1.55,lane['label'],color)
        steps=lane['steps']; cw=(9.72-.22*(len(steps)-1))/len(steps)
        for j,c in enumerate(steps):
            x=2.72+j*(cw+.22)
            box(s,x,y+.19,cw,1.10,fill=PALE if k==0 else 'E8EFF8')
            text(s,x+.09,y+.31,cw-.18,.37,c[0],15.5,NAVY,True)
            text(s,x+.09,y+.77,cw-.18,.38,c[1],11.5,MUTED)
            if j<len(steps)-1: line(s,x+cw,y+.72,x+cw+.20,y+.72,color,1.3,True)
    text(s,.81,5.94,11.7,.34,d.get('below',''),15,MUTED)

def render_levels(s,d):
    y=2.26
    for k,row in enumerate(d['rows']):
        yy=y+k*.72
        box(s,.56,yy,12.18,.63,fill=WHITE,line=LINE)
        chip(s,.72,yy+.14,.79,row[0],[MUTED,TEAL,BLUE,PURPLE,ORANGE][k])
        text(s,1.75,yy+.15,2.24,.40,row[1],16.5,NAVY,True)
        text(s,4.10,yy+.15,8.38,.40,row[2],15,INK)

def render_bars(s,d):
    for j,metric in enumerate(d['metrics']):
        x=.56+j*4.12
        box(s,x,2.26,3.94,3.91,fill=WHITE,line=LINE)
        text(s,x+.19,2.48,3.55,.39,metric['title'],18,NAVY,True)
        data=CategoryChartData(); data.categories=metric['categories']
        data.add_series(metric['title'],metric['values'])
        chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(x+.18),Inches(3.05),
                               Inches(3.56),Inches(2.25),data).chart
        chart.has_legend=False; chart.has_title=False
        chart.chart_style=10
        chart.value_axis.minimum_scale=0; chart.value_axis.maximum_scale=metric['max']
        chart.value_axis.has_major_gridlines=False
        chart.value_axis.tick_labels.font.size=Pt(10)
        chart.category_axis.tick_labels.font.size=Pt(11)
        chart.category_axis.tick_labels.font.name=FONT
        plot=chart.plots[0]; plot.gap_width=65
        plot.has_data_labels=True; plot.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.font.size=Pt(16); plot.data_labels.font.bold=True
        plot.data_labels.number_format=metric.get('format','0')
        for n,p in enumerate(chart.series[0].points):
            p.format.fill.solid(); p.format.fill.fore_color.rgb=rgb(MUTED if n==0 else TEAL)
            p.format.line.fill.background()
        text(s,x+.19,5.55,3.52,.39,metric['delta'],17,metric.get('color',TEAL),True)

def render_search(s,d):
    box(s,.56,2.26,7.20,3.90,fill=WHITE,line=LINE)
    chip(s,.83,2.47,1.91,'训练反馈 → 候选')
    nodes=[(1.05,3.14,2.00,.88,'父策略','固定版本',TEAL),
           (4.18,2.97,2.40,.88,'反思变异','经验 / 动作 / loop',BLUE),
           (4.18,4.17,2.40,.88,'多父重组','互补分支联合',PURPLE),
           (1.05,5.14,2.00,.67,'Frontier','保留互补候选',ORANGE)]
    for x,y,w,h,a,b,c in nodes:
        box(s,x,y,w,h,fill=BG,line=LINE)
        text(s,x+.12,y+.10,w-.24,.28,a,15,c,True)
        text(s,x+.12,y+.46,w-.24,.22,b,10.5,MUTED)
    line(s,3.05,3.59,4.12,3.40,BLUE,1.7,True)
    line(s,3.05,3.59,4.12,4.60,PURPLE,1.7,True)
    line(s,6.61,3.42,7.28,3.42,TEAL,1.7,True)
    line(s,7.28,3.42,7.28,5.50,TEAL,1.7)
    line(s,7.28,5.50,3.09,5.50,TEAL,1.7,True)
    line(s,2.05,5.13,2.05,4.07,ORANGE,1.7,True)
    card(s,8.00,2.26,4.74,3.90,'每一步都要回答',d['body'],TEAL)

def render_cover(s,d):
    s.background.fill.fore_color.rgb=rgb(NAVY)
    box(s,9.03,0,4.30,H,fill='1B3B59',radius=False)
    text(s,.76,.65,8.10,.37,'导师阶段汇报  /  AGENT HARNESS',13,'77D1C7',True)
    text(s,.76,1.69,8.20,1.75,d['title'],36,WHITE,True)
    text(s,.79,3.68,7.86,1.00,d['subtitle'],21,'CFDFEC')
    text(s,.79,5.77,7.84,.63,'2026.10.09  ·  约 30 分钟\n代码快照 f44bbaf4 · 现存实验产物核查',13,'B0C5D6')
    for k,(a,b) in enumerate([('运行','C++ / Python Harness'),('学习','轨迹 → 经验 → 策略'),('验证','同题对照 / 真实应用')]):
        yy=1.50+k*1.55
        box(s,9.53,yy,3.04,1.16,fill='254A68')
        text(s,9.75,yy+.20,2.60,.35,a,21,'77D1C7',True)
        text(s,9.75,yy+.71,2.60,.30,b,11,WHITE)

for i,d in enumerate(DATA['slides'],1):
    s=base(d,i)
    kind=d.get('type','cards')
    if kind=='cover': render_cover(s,d)
    elif kind=='cards': render_cards(s,d)
    elif kind=='flow': render_flow(s,d)
    elif kind=='lanes': render_lanes(s,d)
    elif kind=='levels': render_levels(s,d)
    elif kind=='table': table(s,d['headers'],d['rows'],d.get('widths'),h=d.get('table_h',3.80),font=d.get('table_font',15))
    elif kind=='bars': render_bars(s,d)
    elif kind=='search': render_search(s,d)
    else: raise ValueError(kind)
    if d.get('takeaway'): takeaway(s,d['takeaway'],d.get('takeaway_color',TEAL))
    notes=f"第{i}页｜{d['title']}\n建议用时：{d.get('seconds',60)}秒\n\n{d.get('notes','')}\n\n来源与证据边界：\n"+'\n'.join(d.get('sources',[]))
    s.notes_slide.notes_text_frame.text=notes

prs.save(OUT)
notes=['# AgentFlow 导师汇报逐页讲稿','',
       '约30分钟。主讲页1—30；31页起为答疑附录。依据当前源码、现存结果摘要和明确标注的历史状态文档，不代表本轮重跑实验。','']
for i,d in enumerate(DATA['slides'],1):
    notes.extend([f"## {i:02d}. {d['title'].replace(chr(10),' ')}",'',
                  f"建议用时：{d.get('seconds',60)}秒；{'主讲' if i<=30 else '答疑附录'}。",'',d.get('notes',''),'',
                  '**来源：**','']+[f'- {x}' for x in d.get('sources',[])]+[''])
(ROOT/'逐页讲稿.md').write_text('\n'.join(notes),encoding='utf-8')
repo=ROOT.parents[1]
idx=['# 逐页证据定位','','代码快照为f44bbaf4；JSON行号1表示压缩存储，字段路径见备注。','']
import re
for i,d in enumerate(DATA['slides'],1):
    idx.extend([f"## {i:02d} {d['title'].replace(chr(10),' ')}",''])
    for src in d['sources']:
        m=re.match(r'((?:tools|packages|frontend|src|trace_output)/[^:；（\s]+|CMakeLists\.txt)(?::(\d+))?',src)
        if m and (repo/m[1]).exists():
            idx.append(f'- [{m[1]}](../../{m[1]})'+('；'+src[len(m[1]):] if len(src)>len(m[1]) else ''))
        else: idx.append('- '+src)
    idx.append('')
(ROOT/'源码与产物定位.md').write_text('\n'.join(idx),encoding='utf-8')
print(json.dumps({'output':str(OUT),'slides':len(prs.slides),
                  'main_seconds':sum(d.get('seconds',60) for d in DATA['slides'][:30])},ensure_ascii=False))
