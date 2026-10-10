#!/usr/bin/env python3
"""Validate the generated presentation and its PDF, without running the repository."""
from pathlib import Path
import json,re,sys,unicodedata
if Path('/tmp/agentflow-report-pydeps').exists():sys.path.insert(0,'/tmp/agentflow-report-pydeps')
from pptx import Presentation
import fitz

ROOT=Path(__file__).resolve().parent
D=json.loads((ROOT/'deck_content.json').read_text())
deck=Presentation(ROOT/'AgentFlow_自进化系统_导师汇报_20261009.pptx')
pdf=fitz.open(ROOT/'AgentFlow_自进化系统_导师汇报_20261009.pdf')
checks=[]
def check(name,condition,detail):
    if not condition:raise AssertionError(f'{name}: {detail}')
    checks.append(dict(name=name,status='pass',detail=detail))
def norm(t):return re.sub(r'\s+','',unicodedata.normalize('NFKC',t))
check('Slide and PDF page count',len(deck.slides)==len(pdf)==len(D['slides'])==38,'30 main slides and 8 appendix slides')
check('Main talk timing',sum(s['seconds'] for s in D['slides'][:30])==1800,'1800 seconds')
check('Speaker notes and sources',all(s.notes_slide.notes_text_frame.text.strip() and d['sources'] for s,d in zip(deck.slides,D['slides'])),'38 pages contain notes and provenance')
check('Editable charts',sum(z.has_chart for s in deck.slides for z in s.shapes)==3,'Three native charts with editable workbook data')
outside=[]
for i,s in enumerate(deck.slides,1):
    for z in s.shapes:
        if z.left<0 or z.top<0 or z.left+z.width>deck.slide_width+10000 or z.top+z.height>deck.slide_height+10000:
            outside.append([i,z.name])
check('Shapes within slide canvas',not outside,outside)
missing=[]
for i,(s,p) in enumerate(zip(deck.slides,pdf),1):
    t=norm(p.get_text())
    # Check every substantive body paragraph survives PDF rendering.
    for z in s.shapes:
        if not z.has_text_frame or z.top>6.95*914400:continue
        for q in z.text_frame.paragraphs:
            v=norm(q.text)
            if len(v)>2 and v not in t:missing.append([i,q.text])
check('Rendered substantive text',not missing,missing)
check('PDF pages have expected aspect ratio',all(abs(p.rect.width/p.rect.height-16/9)<.01 for p in pdf),'16:9')
check('Evidence qualification explicit','best_available_not_qualified' in (ROOT/'逐页讲稿.md').read_text() and 'independent_holdout=false' in (ROOT/'逐页讲稿.md').read_text(),'Delivery and development-set limits retained')
receipt=dict(validation='Presentation structure, native charts, sources, notes, page layout and rendered text; does not test repository software or reproduce experiments',passed=len(checks),checks=checks)
(ROOT/'文稿校验回执.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2))
print(f'{len(checks)}/{len(checks)} presentation checks PASS')
